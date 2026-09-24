import io
import uuid
from types import SimpleNamespace

import pytest
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from rest_framework.test import APIClient
from docx import Document as WordDocument

from apps.applications.models import Application
from apps.enterprise.models import Membership
from apps.knowledge.execution import execute_knowledge_index
from apps.knowledge.extractors import extract_document
from apps.knowledge.models import KnowledgeBase, KnowledgeChunk, KnowledgeDocument
from apps.knowledge.index_backend import replace_document_index, vector_candidates
from apps.knowledge.retrieval import search, lexical_text
from modules.execution.models import Run, RunArtifact
from modules.execution.infrastructure.artifacts import get_artifact_storage, artifact_token
from app_center.research_assistant import runtime
from ..content import SECTIONS, validate_output, markdown, document_content
from ..models import ResearchProject, ResearchSource, ResearchResult


class Sink:
    cancelled = False

    def __init__(self):
        self.events = []

    def emit(self, kind, value):
        self.events.append((kind, value))


@pytest.fixture
def ctx(db, settings, tmp_path):
    settings.ROOT_URLCONF = "app_center.research_assistant.backend.tests.urls"
    settings.ARTIFACT_ROOT = tmp_path / "artifacts"
    settings.ARTIFACT_STORAGE_BACKEND = "local"
    settings.MY_DRIVE_ROOT = tmp_path / "drive"
    owner = get_user_model().objects.create_user(username="research-owner")
    other = get_user_model().objects.create_user(username="research-other")
    org = owner.owned_organizations.get()
    Membership.objects.create(organization=org, user=other, role=Membership.Role.ADMIN)
    call_command("sync_app_center", package_id="research-assistant", organization_id=str(org.id))
    app = Application.objects.get(organization=org, slug="research-assistant")
    app.visibility = "organization"
    app.save(update_fields=["visibility"])
    client = APIClient()
    client.force_authenticate(owner)
    root = f"/api/v1/organizations/{org.id}"
    url = f"{root}/applications/{app.id}/research-assistant"
    response = client.post(url + "/projects", {"title": "行业研究", "objective": "比较价格与增长观点"}, format="json")
    assert response.status_code == 201, response.data
    project = ResearchProject.objects.get(pk=response.data["id"])
    return SimpleNamespace(owner=owner, other=other, org=org, app=app, client=client, root=root, url=url,
                           project=project, project_url=f"{url}/projects/{project.id}")


def add(ctx, text="价格上涨，但销量下降。", title="分析报告"):
    response = ctx.client.post(ctx.project_url + "/sources", {"title": title, "text": text}, format="json")
    assert response.status_code in (200, 202), response.data
    return ResearchSource.objects.select_related("document").get(pk=response.data["id"])


def index(source):
    run = source.document.indexing_run
    execute_knowledge_index({"run_id": str(run.id), "definition_snapshot": run.definition_snapshot}, Sink())
    source.document.refresh_from_db()
    assert source.document.status == "ready"


def generate(ctx, sources, kind="report", key="generate-1"):
    response = ctx.client.post(ctx.project_url + "/results", {"source_ids": [str(s.id) for s in sources],
        "kind": kind, "instruction": "说明共识和分歧"}, format="json", HTTP_IDEMPOTENCY_KEY=key)
    assert response.status_code in (200, 202), response.data
    return ResearchResult.objects.select_related("run", "project").get(pk=response.data["id"])


def fake_model(project, prompt, data, config):
    if "片段" in data:
        return {"evidence": [{"chunk_id": c["chunk_id"], "quote": c["text"][:100], "note": "资料观点"} for c in data["片段"][:6]]}
    ids = list(data["证据"])
    sections = [{"heading": name, "items": [{"type": "fact" if ids else "gap", "text": "资料显示价格存在变化。" if ids else "资料不足。", "evidence_ids": ids[:2]}]} for name in data["章节"]]
    if ids:
        for section in sections:
            if section["heading"] == "原文摘录":
                section["items"] = [{"type": "quote", "text": data["证据"][ids[0]]["quote"], "evidence_ids": ids[:1]}]
    if data["章节"] == SECTIONS["comparison"]:
        sections[0]["items"] = [{"source_id": s["source_id"], "type": "fact", "text": "资料观点。",
            "evidence_ids": [key for key in ids if data["证据"][key]["source_id"] == s["source_id"]][:1]} for s in data["资料覆盖"]]
    return {"sections": sections}


def finish(ctx, result, monkeypatch):
    monkeypatch.setattr(runtime, "call_model", fake_model)
    output = runtime.execute({"run_id": str(result.run_id), "organization_id": str(ctx.org.id), "input": result.run.input}, Sink())
    result.run.output_summary = output
    result.run.status = "succeeded"
    result.run.save(update_fields=["output_summary", "status"])
    return output


def test_project_private_through_every_knowledge_and_run_entry(ctx):
    source = add(ctx)
    index(source)
    doc = source.document
    assert doc.knowledge_base.scope == "research"
    assert search(organization_id=ctx.org.id, query="价格", knowledge_base_ids=[doc.knowledge_base_id])[1] == []
    assert search(organization_id=ctx.org.id, query="价格", knowledge_base_ids=[doc.knowledge_base_id], internal=True)[1]
    for user in (ctx.owner, ctx.other):
        ctx.client.force_authenticate(user)
        listed = ctx.client.get(ctx.root + "/knowledge-bases/")
        assert not any(item["id"] == doc.knowledge_base_id for item in listed.data)
        base = f"{ctx.root}/knowledge-bases/{doc.knowledge_base_id}"
        for suffix in ("/", "/documents/", f"/documents/{doc.id}/", f"/documents/{doc.id}/content/"):
            assert ctx.client.get(base + suffix).status_code == 404
        assert ctx.client.post(base + f"/documents/{doc.id}/reindex/").status_code == 404
        assert ctx.client.post(ctx.root + "/knowledge-search/", {"query": "价格", "knowledge_base_ids": [doc.knowledge_base_id]}, format="json").status_code == 404
    assert ctx.client.get(ctx.project_url).status_code == 404
    assert ctx.client.get(ctx.project_url + f"/sources/{source.id}/content").status_code == 404
    run_url = f"{ctx.root}/runs/{doc.indexing_run_id}"
    for suffix in ("", "/events", "/attempts", "/artifacts", "/snapshot", "/children", "/stream"):
        assert ctx.client.get(run_url + suffix).status_code == 404, suffix
    assert ctx.client.post(run_url + "/commands", {"type": "cancel", "idempotency_key": "deny"}, format="json").status_code == 404
    assert ctx.client.delete(run_url).status_code == 404
    assert ctx.client.delete(base + f"/documents/{doc.id}/").status_code == 204
    doc.refresh_from_db()
    assert not doc.is_deleted


@pytest.mark.parametrize("kind", ["report", "comparison", "writing_pack"])
def test_three_outputs_history_citations_and_exports(ctx, monkeypatch, kind):
    left, right = add(ctx), add(ctx, "价格下降，增长没有出现。", "另一份观点")
    index(left); index(right)
    result = generate(ctx, [left, right], kind=kind)
    replay = generate(ctx, [left, right], kind=kind)
    assert replay.id == result.id
    output = finish(ctx, result, monkeypatch)
    assert [s["heading"] for s in output["sections"]] == SECTIONS[kind]
    assert len(output["coverage"]) == 2 and all(c["chunks_reviewed"] == c["chunks_total"] for c in output["coverage"])
    url = ctx.project_url + f"/results/{result.id}"
    assert ctx.client.get(url).data["output"]["citations"]
    first = output["citations"][0]
    citation_url = url + f"/citations/{first['id']}"
    citation = ctx.client.get(citation_url)
    assert citation.status_code == 200 and first["quote"] in citation.data["context"]
    ctx.client.delete(ctx.project_url + f"/sources/{left.id}")
    assert ctx.client.get(citation_url).status_code == 200
    raw = ctx.client.get(url + "/download").content.decode()
    assert "参考资料" in raw and first["quote"] in raw and "citation=" in raw
    from app_center.documents.backend.content import validate_content, plain_text
    content = validate_content(document_content(output, "https://example.com", ctx.project, result))
    assert first["quote"] in plain_text(content)
    assert ctx.client.get(ctx.project_url + "/results").data["count"] == 1
    ctx.client.force_authenticate(ctx.other)
    assert ctx.client.get(url).status_code == 404
    assert ctx.client.get(citation_url).status_code == 404


def test_forged_quotes_fail_without_publishing(ctx, monkeypatch):
    source = add(ctx); index(source)
    result = generate(ctx, [source])
    monkeypatch.setattr(runtime, "call_model", lambda *args: {"evidence": [{"chunk_id": source.document.chunks.first().id, "quote": "虚构数据", "note": "错误"}]})
    with pytest.raises(RuntimeError, match="原文证据校验"):
        runtime.execute({"run_id": str(result.run_id), "organization_id": str(ctx.org.id), "input": result.run.input}, Sink())
    assert ctx.client.get(ctx.project_url + f"/results/{result.id}").data["output"] is None


def test_unknown_citation_and_uncited_fact_rejected():
    value = {"sections": [{"heading": heading, "items": [{"type": "fact", "text": "无证据结论", "evidence_ids": ["fake"]}]} for heading in SECTIONS["report"]]}
    with pytest.raises(ValueError, match="不属于"):
        validate_output(value, {}, "report")
    value["sections"][0]["items"][0]["evidence_ids"] = []
    with pytest.raises(ValueError, match="缺少"):
        validate_output(value, {}, "report")


def test_limits_duplicates_and_cancel(ctx, settings):
    settings.RESEARCH_MAX_SOURCES = 1
    source = add(ctx)
    assert add(ctx).id == source.id
    assert ctx.client.post(ctx.project_url + "/sources", {"text": "第二份资料"}, format="json").status_code == 400
    index(source)
    result = generate(ctx, [source])
    response = ctx.client.post(ctx.project_url + f"/results/{result.id}/cancel")
    assert response.status_code == 200
    result.run.refresh_from_db()
    assert result.run.status == "cancelled"
    response = ctx.client.post(ctx.project_url + "/results", {"kind": "writing_pack", "source_ids": [str(source.id)]}, format="json", HTTP_IDEMPOTENCY_KEY="generate-1")
    assert response.status_code == 409


def test_docx_tables_keep_order_and_positions():
    document = WordDocument()
    document.add_heading("价格", 1)
    document.add_paragraph("表格之前")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "产品"
    table.cell(0, 1).text = "100 元"
    document.add_paragraph("表格之后")
    stream = io.BytesIO(); document.save(stream)
    sections = extract_document(stream.getvalue(), "价格.docx")
    assert [s.text for s in sections] == ["价格", "表格之前", "产品 | 100 元", "表格之后"]
    assert sections[2].paragraph_number == 3 and sections[2].section_path == ("价格",)


def test_document_filter_before_candidate_limit(ctx):
    target = add(ctx); index(target)
    for i in range(55):
        document = KnowledgeDocument.objects.create(organization=ctx.org, knowledge_base=ctx.project.knowledge_base,
            title=f"干扰{i}", status="ready", active_revision=1)
        chunk = KnowledgeChunk.objects.create(organization=ctx.org, document=document, revision=1, position=0,
            content="价格", lexical_text=lexical_text("价格"))
        replace_document_index(document, [chunk])
    _, rows = search(organization_id=ctx.org.id, knowledge_base_ids=[ctx.project.knowledge_base_id],
                     query="价格", document_ids=[target.document_id], internal=True)
    assert [row["document_id"] for row in rows] == [target.document_id]


def test_delete_project_denies_existing_citations_and_run(ctx, monkeypatch):
    source = add(ctx); index(source)
    result = generate(ctx, [source]); finish(ctx, result, monkeypatch)
    key = source.document.source_object_key
    response = ctx.client.delete(ctx.project_url)
    assert response.status_code == 204
    assert ctx.client.get(ctx.project_url).status_code == 404
    assert ctx.client.get(f"{ctx.root}/runs/{result.run_id}").status_code == 404
    assert ctx.client.get(ctx.project_url + f"/sources/{source.id}/content").status_code == 404
    assert not source.document.chunks.exists()
    from modules.execution.infrastructure.artifacts import ArtifactObjectUnavailable
    with pytest.raises(ArtifactObjectUnavailable):
        get_artifact_storage().open(key)
    assert ctx.client.delete(ctx.project_url).status_code == 204


def install(ctx, package):
    call_command("sync_app_center", package_id=package, organization_id=str(ctx.org.id))
    return Application.objects.get(organization=ctx.org, slug=package)


def test_online_document_snapshot_and_idempotent_save(ctx, monkeypatch):
    from app_center.documents.backend.models import Document, DocumentGrant
    app = install(ctx, "documents")
    doc = Document.objects.create(organization=ctx.org, application=app, owner=ctx.other, title="共享资料",
        content={"type": "doc", "content": [{"type": "heading", "attrs": {"level": 1}, "content": [{"type": "text", "text": "增长观点"}]},
        {"type": "paragraph", "content": [{"type": "text", "text": "价格上涨，但销量下降。"}]}]})
    body = {"origin_type": "document", "origin_id": str(doc.id), "application_id": app.id}
    assert ctx.client.post(ctx.project_url + "/sources", body, format="json").status_code == 404
    DocumentGrant.objects.create(document=doc, user=ctx.owner, role="viewer")
    response = ctx.client.post(ctx.project_url + "/sources", body, format="json")
    assert response.status_code == 202, response.data
    source = ResearchSource.objects.select_related("document").get(pk=response.data["id"])
    index(source)
    result = generate(ctx, [source]); output = finish(ctx, result, monkeypatch)
    doc.delete()
    citation = output["citations"][0]
    assert ctx.client.get(ctx.project_url + f"/results/{result.id}/citations/{citation['id']}").status_code == 200
    url = ctx.project_url + f"/results/{result.id}/export"
    first = ctx.client.post(url, {"target": "document", "application_id": app.id}, format="json", HTTP_IDEMPOTENCY_KEY="save-doc")
    second = ctx.client.post(url, {"target": "document", "application_id": app.id}, format="json", HTTP_IDEMPOTENCY_KEY="save-doc")
    assert first.status_code == 200, first.data
    assert first.data == second.data
    saved = Document.objects.get(pk=first.data["id"])
    assert saved.owner == ctx.owner and not saved.grants.exists()
    assert citation["quote"] in saved.plain_text


def test_drive_export_import_and_quota(ctx, monkeypatch, settings):
    app = install(ctx, "my-drive")
    source = add(ctx); index(source)
    result = generate(ctx, [source]); finish(ctx, result, monkeypatch)
    export_url = ctx.project_url + f"/results/{result.id}/export"
    body = {"target": "drive", "application_id": app.id}
    settings.MY_DRIVE_QUOTA_BYTES = 1
    assert ctx.client.post(export_url, body, format="json", HTTP_IDEMPOTENCY_KEY="drive-small").status_code == 400
    settings.MY_DRIVE_QUOTA_BYTES = 1024 * 1024
    first = ctx.client.post(export_url, body, format="json", HTTP_IDEMPOTENCY_KEY="drive-save")
    assert first.status_code == 200, first.data
    second = ctx.client.post(export_url, body, format="json", HTTP_IDEMPOTENCY_KEY="drive-save")
    assert first.data == second.data
    response = ctx.client.post(ctx.project_url + "/sources", {"origin_type": "drive", "application_id": app.id, "origin_id": first.data["id"]}, format="json")
    assert response.status_code == 202, response.data
    from app_center.my_drive.backend.models import DriveEntry
    entry = DriveEntry.objects.get(pk=first.data["id"])
    import uuid
    entry.trash_batch = uuid.uuid4(); entry.save(update_fields=["trash_batch"])
    response = ctx.client.post(ctx.project_url + "/sources", {"origin_type": "drive", "application_id": app.id, "origin_id": first.data["id"]}, format="json")
    assert response.status_code == 404


def test_private_artifact_tokens_require_owner_and_live_project(ctx):
    source = add(ctx)
    run = source.document.indexing_run
    get_artifact_storage().put("research-test/source.txt", b"private")
    artifact = RunArtifact.objects.create(organization=ctx.org, run=run, kind="test", object_key="research-test/source.txt",
                                         content_hash="abc", mime_type="text/plain", size=7)
    url = f"{ctx.root}/runs/{run.id}/artifacts/{artifact.id}/content?token={artifact_token(artifact)}"
    assert ctx.client.get(url).status_code == 200
    ctx.client.force_authenticate(ctx.other)
    assert ctx.client.get(url).status_code == 404
    ctx.client.force_authenticate(None)
    assert ctx.client.get(url).status_code == 404


def test_failure_limits_and_invalid_file_types(ctx, settings):
    settings.RESEARCH_MAX_TEXT_BYTES = 3
    response = ctx.client.post(ctx.project_url + "/sources", {"text": "超过上限"}, format="json")
    assert response.status_code == 400
    response = ctx.client.post(ctx.project_url + "/sources", {"file": SimpleUploadedFile("old.doc", b"not-docx")}, format="multipart")
    assert response.status_code == 400
    response = ctx.client.post(ctx.project_url + "/sources", {"file": SimpleUploadedFile("invalid.pdf", b"not-pdf")}, format="multipart")
    assert response.status_code == 202
    source = ResearchSource.objects.select_related("document").get(pk=response.data["id"])
    run = source.document.indexing_run
    with pytest.raises(Exception):
        execute_knowledge_index({"run_id": str(run.id), "definition_snapshot": run.definition_snapshot}, Sink())
    source.document.refresh_from_db()
    assert source.document.status == "failed"
    assert ctx.client.post(ctx.project_url + f"/sources/{source.id}/retry").status_code == 202


def test_long_sources_review_every_chunk_with_bounded_batches(ctx, monkeypatch):
    source = add(ctx, ("价格变化应结合需求和供应两方面分析。" * 2200))
    index(source)
    result = generate(ctx, [source])
    seen, batch_sizes = [], []

    def model(project, prompt, data, config):
        if "片段" in data:
            seen.extend(c["chunk_id"] for c in data["片段"])
            batch_sizes.append(sum(len(c["text"]) for c in data["片段"]))
        return fake_model(project, prompt, data, config)

    monkeypatch.setattr(runtime, "call_model", model)
    output = runtime.execute({"run_id": str(result.run_id), "organization_id": str(ctx.org.id), "input": result.run.input}, Sink())
    assert set(seen) == set(source.document.chunks.values_list("id", flat=True))
    assert len(batch_sizes) > 1 and max(batch_sizes) <= 18000
    assert output["coverage"][0]["chunks_reviewed"] == len(seen)


def test_insufficient_evidence_is_a_gap_not_a_fact(ctx, monkeypatch):
    ctx.project.objective = "unrelatedtopic"
    ctx.project.save(update_fields=["objective"])
    source = add(ctx, "不相关的背景材料。")
    index(source)
    result = generate(ctx, [source])
    result.instruction = ""
    result.save(update_fields=["instruction"])
    monkeypatch.setattr(runtime, "call_model", lambda project, prompt, data, config:
        {"evidence": []} if "片段" in data else fake_model(project, prompt, data, config))
    output = runtime.execute({"run_id": str(result.run_id), "organization_id": str(ctx.org.id), "input": result.run.input}, Sink())
    assert output["citations"] == []
    assert all(i["type"] == "gap" for s in output["sections"] for i in s["items"])


def test_model_failure_and_cancellation_do_not_publish(ctx, monkeypatch):
    source = add(ctx); index(source)
    result = generate(ctx, [source])
    sink = Sink()

    def model(*args):
        raise RuntimeError("模型暂时不可用")

    monkeypatch.setattr(runtime, "call_model", model)
    with pytest.raises(RuntimeError, match="模型暂时不可用"):
        runtime.execute({"run_id": str(result.run_id), "organization_id": str(ctx.org.id), "input": result.run.input}, sink)
    sink.cancelled = True
    assert runtime.execute({"run_id": str(result.run_id), "organization_id": str(ctx.org.id), "input": result.run.input}, sink) == {}
    assert ctx.client.get(ctx.project_url + f"/results/{result.id}").data["output"] is None


def test_comparison_cannot_silently_omit_a_source():
    evidence = {"e1": {"quote": "原文", "source_id": "s1"}}
    value = {"sections": [{"heading": heading, "items": [{"source_id": "s1", "type": "fact", "text": "观点", "evidence_ids": ["e1"]}]} for heading in SECTIONS["comparison"]]}
    with pytest.raises(ValueError, match="每份资料"):
        validate_output(value, evidence, "comparison", [{"source_id": "s1", "title": "A"}, {"source_id": "s2", "title": "B"}])


def test_cross_organization_and_membership_revocation(ctx):
    source = add(ctx)
    other_org = ctx.other.owned_organizations.get()
    assert ctx.client.get(ctx.project_url.replace(str(ctx.org.id), str(other_org.id))).status_code in (403, 404)
    Membership.objects.filter(organization=ctx.org, user=ctx.owner).update(is_active=False)
    assert ctx.client.get(ctx.project_url).status_code == 403
    assert ctx.client.get(f"{ctx.root}/runs/{source.document.indexing_run_id}").status_code == 403


def test_cleanup_retry_and_running_cancel_are_safe(ctx, monkeypatch):
    source = add(ctx); index(source)
    result = generate(ctx, [source])
    result.run.status = "cancelling"
    result.run.save(update_fields=["status"])
    from .. import services
    real_delete = services.delete_artifact_object
    monkeypatch.setattr(services, "delete_artifact_object", lambda key: (_ for _ in ()).throw(OSError("storage offline")))
    assert ctx.client.delete(ctx.project_url).status_code == 503
    ctx.project.refresh_from_db()
    assert ctx.project.deleted_at and ctx.project.cleanup_pending
    assert ctx.client.get(ctx.project_url).status_code == 404
    monkeypatch.setattr(services, "delete_artifact_object", real_delete)
    call_command("cleanup_research_projects")
    ctx.project.refresh_from_db()
    assert not ctx.project.cleanup_pending


def test_vector_candidates_filter_document_before_limit(ctx, settings):
    source = add(ctx); index(source)
    chunk = source.document.chunks.first()
    vector = [1.0] + [0.0] * (settings.KNOWLEDGE_EMBEDDING_DIMENSIONS - 1)
    chunk.embedding = vector
    replace_document_index(source.document, [chunk])
    assert vector_candidates(ctx.org.id, [ctx.project.knowledge_base_id], vector, 1, document_ids=[source.document_id])[0][0] == chunk.id
    assert vector_candidates(ctx.org.id, [ctx.project.knowledge_base_id], vector, 1, document_ids=[999999]) == []


def test_export_links_use_frontend_origin_and_reject_untrusted_host(settings):
    from django.test import RequestFactory
    from ..services import public_origin
    settings.RESEARCH_PUBLIC_URL = ""
    settings.ALLOWED_HOSTS = ["testserver", "localhost"]
    settings.CORS_ALLOWED_ORIGINS = []
    request = RequestFactory().get("/api/v1/", HTTP_REFERER="http://localhost:3030/applications/1/research-assistant")
    assert public_origin(request) == "http://localhost:3030"
    request = RequestFactory().get("/api/v1/", HTTP_REFERER="https://untrusted.invalid/")
    assert public_origin(request) == "http://testserver"
    settings.RESEARCH_PUBLIC_URL = "https://studio.example.com"
    assert public_origin(request) == "https://studio.example.com"
