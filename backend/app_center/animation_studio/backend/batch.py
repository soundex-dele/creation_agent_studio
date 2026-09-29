import copy
import csv
import io
import json
import math
import zipfile

from django.db import transaction
from django.shortcuts import get_object_or_404
from rest_framework.exceptions import ValidationError
from rest_framework.response import Response
from modules.execution.models import Run
from modules.execution.application.runs import create_run
from modules.execution.application.commands import submit_run_command
from modules.execution.infrastructure.artifacts import open_artifact
from .views import BaseView
from .models import AnimationBatch, AnimationPreset, AnimationProject, AnimationVersion
from .documents import builtins, validate_document
from .projects import document_assets
from .studio_validation import export_options

TERMINAL = {"succeeded", "failed", "cancelled"}


def validate_fields(template):
    fields = template.get("fields", [])
    scenes = {s["id"]: s for s in template["document"]["scenes"]}
    names = set()
    if not isinstance(fields, list) or len(fields) > 100: raise ValidationError("模板字段最多100个。")
    for field in fields:
        if not isinstance(field, dict) or not isinstance(field.get("name"), str) or not field["name"] or field["name"] in names: raise ValidationError("模板字段名无效或重复。")
        names.add(field["name"])
        item = scenes.get(field.get("scene_id"))
        if not item or item.get("locked") or field.get("property") not in ("title", "body", "narration", "assets") or field.get("type") not in ("text", "number", "image"):
            raise ValidationError("模板字段必须指向未锁定场景中的文字、数字或图片。")
        if (field["type"] == "image") != (field["property"] == "assets"): raise ValidationError("图片字段必须映射至素材。")


def mapped_document(template, row):
    if not isinstance(row, dict): raise ValidationError("每行数据须为对象。")
    doc = copy.deepcopy(template["document"])
    for field in template["fields"]:
        value = row.get(field["name"], "")
        if value is None or value == "": raise ValidationError(f"缺少字段 {field['name']}。")
        if field["type"] == "number":
            try:
                if not math.isfinite(float(value)): raise ValueError()
            except (ValueError, TypeError): raise ValidationError(f"字段 {field['name']} 须为数字。")
        item = next(s for s in doc["scenes"] if s["id"] == field["scene_id"])
        item[field["property"]] = [str(value)] if field["type"] == "image" else str(value)
    return validate_document(doc, complete=True)


def template_for(user, app, identifier):
    builtin = next((p for p in builtins() if p["id"] == identifier), None)
    if builtin: return builtin["data"]
    return get_object_or_404(AnimationPreset.objects.for_organization(app.organization_id), application=app, owner=user, pk=identifier, kind="template").data


class BatchImportView(BaseView):
    def post(self, request, **kwargs):
        self.application(); upload = request.FILES.get("file")
        if not upload or upload.size > 10 * 1024 * 1024: raise ValidationError("请上传不超过10 MB的 CSV 或 XLSX。")
        sheets = []
        try:
            if upload.name.lower().endswith(".csv"):
                rows = list(csv.reader(io.StringIO(upload.read().decode("utf-8-sig"))))
                sheets = [{"name": "CSV", "rows": rows}]
            elif upload.name.lower().endswith(".xlsx"):
                from openpyxl import load_workbook
                content = upload.read()
                with zipfile.ZipFile(io.BytesIO(content)) as archive:
                    if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024: raise ValidationError("工作簿解压后过大。")
                workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
                try:
                    if len(workbook.worksheets) > 20: raise ValidationError("工作表最多20个。")
                    for sheet in workbook.worksheets:
                        rows = []
                        for row in sheet.iter_rows():
                            if len(rows) > 100 or len(row) > 100: raise ValidationError("每表最多100行数据、100列。")
                            if any(c.data_type == "f" for c in row): raise ValidationError("请把公式转换为值后导入。")
                            rows.append([str(c.value) if c.value is not None else "" for c in row])
                        sheets.append({"name": sheet.title, "rows": rows})
                finally: workbook.close()
            else: raise ValidationError("只支持 CSV 和 XLSX。")
        except (UnicodeError, ValueError, zipfile.BadZipFile, OSError):
            raise ValidationError("文件无法读取，请检查格式。")
        for sheet in sheets:
            rows = sheet.pop("rows")
            if not rows or len(rows) > 101 or any(len(r) > 100 for r in rows): raise ValidationError("每表须有表头，最多100行数据、100列。")
            headers = [str(v).strip() for v in rows[0]]
            if len(set(headers)) != len(headers) or any(not h for h in headers): raise ValidationError("列名不能为空或重复。")
            sheet["columns"] = headers
            sheet["rows"] = [dict(zip(headers, r)) for r in rows[1:]]
        return Response({"sheets": sheets})


class BatchesView(BaseView):
    def get(self, request, **kwargs):
        app = self.application()
        return Response({"results": [{"id": str(b.id), "created_at": b.created_at.isoformat(), "status": b.run.status if b.run else "queued", "count": len(b.rows)} for b in AnimationBatch.objects.for_organization(app.organization_id).filter(application=app, owner=request.user).select_related("run").order_by("-created_at")[:100]]})

    @transaction.atomic
    def post(self, request, **kwargs):
        app = self.application(); template = template_for(request.user, app, request.data.get("template_id"))
        validate_fields(template)
        rows = request.data.get("rows", [])
        if not isinstance(rows, list) or not 1 <= len(rows) <= 100: raise ValidationError("每批须为1–100行。")
        prepared, errors = [], []
        for index, row in enumerate(rows):
            try:
                doc = mapped_document(template, row); document_assets(request.user, app, doc)
                prepared.append({"document": doc, "title": str(row.get("title") or f"批量作品 {index+1}")[:200]})
            except ValidationError as exc: errors.append({"row": index + 1, "detail": exc.detail})
        if errors: return Response({"detail": "请修正错误行。", "errors": errors}, status=400)
        if request.data.get("validate_only"): return Response({"valid": True, "count": len(prepared)})
        options = export_options(request.data["export_options"]) if request.data.get("export_options") else {}
        # Stable UUID per idempotency key prevents orphan duplicate batches.
        import uuid
        key = request.headers.get("Idempotency-Key", "")
        if not key or len(key) > 160: raise ValidationError("请提供有效 Idempotency-Key。")
        identifier = uuid.uuid5(uuid.NAMESPACE_URL, f"{app.organization_id}:{app.id}:{request.user.id}:{key}")
        batch, created = AnimationBatch.objects.get_or_create(id=identifier, defaults={"organization": app.organization, "application": app,
            "owner": request.user, "template": template, "rows": prepared, "export_options": options})
        if not created and (batch.template != template or batch.export_options != options or [{k:r[k] for k in ('document','title')} for r in batch.rows] != prepared):
            return Response({"detail": "同一幂等键不能用于不同批量数据。"}, status=409)
        response = self.start(app, {"action": "batch", "batch_id": str(batch.id)})
        if response.status_code < 300:
            batch.run_id = response.data["id"]; batch.save(update_fields=["run"])
            return Response({"id": str(batch.id), "run": response.data}, status=response.status_code)
        transaction.set_rollback(True)
        return response


class BatchView(BaseView):
    def item(self, pk):
        app = self.application()
        return app, get_object_or_404(AnimationBatch.objects.for_organization(app.organization_id), application=app, owner=self.request.user, pk=pk)

    def get(self, request, batch_id, **kwargs):
        app, batch = self.item(batch_id)
        result = []
        for index, row in enumerate(batch.rows):
            rid = row.get("export_run_id") or row.get("generation_run_id")
            run = Run.objects.for_organization(app.organization_id).filter(pk=rid, owner=request.user).first() if rid else None
            result.append({"index": index, "title": row["title"], "project_id": row.get("project_id"), "status": run.status if run else "cancelled" if batch.run and batch.run.status == "cancelled" else "pending", "error": row.get("error") or (run.error_message if run else "")})
        return Response({"id": str(batch.id), "rows": result, "run": self.serialize_run(batch.run) if batch.run else None})

    @transaction.atomic
    def post(self, request, batch_id, **kwargs):
        app, batch = self.item(batch_id)
        operation = request.data.get("operation")
        if operation == "cancel":
            for run in [batch.run, *batch.run.child_runs.all()]:
                if run.status not in TERMINAL:
                    submit_run_command(run_id=run.id, organization_id=app.organization_id, actor=request.user, command_type="cancel", idempotency_key=f"cancel-batch:{batch.id}:{run.id}", payload={})
            return self.get(request, batch_id, **kwargs)
        if operation != "retry" or batch.run.status not in TERMINAL: raise ValidationError("请等待批次结束后重试失败项。")
        rows = copy.deepcopy(batch.rows)
        for row in rows:
            for stage in ("generation", "export"):
                rid = row.get(stage + "_run_id")
                if rid and Run.objects.get(pk=rid).status != "succeeded": row.pop(stage + "_run_id", None)
            row.pop("error", None)
        import uuid
        key = request.headers.get("Idempotency-Key", "")
        if not key or len(key) > 160: raise ValidationError("请提供有效 Idempotency-Key。")
        identifier = uuid.uuid5(uuid.NAMESPACE_URL, f"retry:{batch.id}:{key}")
        clone, _ = AnimationBatch.objects.get_or_create(id=identifier, defaults={"organization": app.organization, "application": app, "owner": request.user,
            "template": batch.template, "export_options": batch.export_options, "rows": rows})
        response = self.start(app, {"action": "batch", "batch_id": str(clone.id)})
        if response.status_code < 300:
            clone.run_id = response.data["id"]; clone.save(update_fields=["run"])
            return Response({"id": str(clone.id), "run": response.data}, status=201)
        transaction.set_rollback(True)
        return response


def execute_batch(payload, sink, root, app, values):
    batch = AnimationBatch.objects.for_organization(app.organization_id).get(id=values["batch_id"], application=app, owner=root.owner)
    with transaction.atomic():
        batch = AnimationBatch.objects.select_for_update().get(pk=batch.id)
        active = []
        for index, row in enumerate(batch.rows):
            if sink.cancelled: return {"cancelled": True}
            if row.get("error"): continue
            generation = Run.objects.filter(pk=row.get("generation_run_id")).first() if row.get("generation_run_id") else None
            exported = Run.objects.filter(pk=row.get("export_run_id")).first() if row.get("export_run_id") else None
            current = exported or generation
            if current and current.status not in TERMINAL:
                active.append(current.id); continue
            if current and current.status in ("failed", "cancelled"):
                row["error"] = current.error_message or current.status; continue
            if generation and (not batch.export_options or exported): continue
            if len(active) >= 2: continue
            if not generation:
                p = AnimationProject.objects.filter(pk=row.get("project_id"), owner=root.owner, application=app).first() if row.get("project_id") else None
                if not p:
                    p = AnimationProject.objects.create(organization=app.organization, application=app, owner=root.owner, title=row["title"], draft=row["document"])
                    row["project_id"] = str(p.id)
                body = {"action": "generate", "project_id": str(p.id), "draft_revision": p.revision, "document": row["document"], "prompt": row["title"], "aspect": row["document"]["aspect"], "style": row["document"]["style"], "asset_ids": [], "duration": 30}
                stage = "generation"
            else:
                body = {"action": "export", "source_run_id": str(generation.id), "export_options": batch.export_options}; stage = "export"
            child = root.child_runs.filter(node_key=f"{index}:{stage}").first()
            if not child:
                child = create_run(organization=app.organization, owner=root.owner, parent=root, node_key=f"{index}:{stage}",
                    executor_kind=root.executor_kind, executor_key=root.executor_key, source_type="application", source_id=str(app.id),
                    definition_snapshot=root.definition_snapshot, input_data=body, priority=root.priority, max_attempts=1, retry_safe=False)
            row[stage + "_run_id"] = str(child.id)
            if stage == "generation": AnimationVersion.objects.get_or_create(run=child, defaults={"organization": app.organization, "project_id": row["project_id"], "document": row["document"], "note": "批量生成"})
            active.append(child.id)
        batch.save(update_fields=["rows"])
    sink.emit("progress.updated", {"stage": "batch", "active": len(active), "count": len(batch.rows)})
    if active:
        sink.wait_for_children(child_run_ids=active, checkpoint={"batch_id": str(batch.id)})
        return {}
    from django.conf import settings
    limit = int(getattr(settings, "EXECUTION_ARTIFACT_MAX_BYTES", 100 * 1024 * 1024))
    output = io.BytesIO(); archive = zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED)
    part = 1; packed = 0; omitted = []
    def publish():
        archive.close()
        sink.create_artifact(kind="animation-batch", filename=f"animations-{part}.zip", content=output.getvalue(),
            mime_type="application/zip", metadata={"count": len(batch.rows), "part": part})
    try:
        for index, row in enumerate(batch.rows):
            if row.get("error"): continue
            run_id = row.get("export_run_id") or row.get("generation_run_id")
            if not run_id: continue
            child = Run.objects.get(pk=run_id)
            kinds = ["animation-video", "animation-gif", "animation-cover", "animation-subtitles"] if row.get("export_run_id") else ["animation-source"]
            for artifact in child.artifacts.filter(kind__in=kinds):
                if sink.cancelled: return {"cancelled": True}
                if artifact.size + 65536 > limit:
                    omitted.append({"row": index + 1, "run_id": str(child.id), "artifact_id": str(artifact.id), "reason": "请在作品中单独下载超大文件"})
                    continue
                if packed and packed + artifact.size + 65536 > limit:
                    publish(); part += 1; output = io.BytesIO(); archive = zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED); packed = 0
                filename = str(artifact.metadata.get("filename") or (artifact.kind + ".bin")).replace("\\", "/").split("/")[-1]
                with open_artifact(artifact) as handle, archive.open(f"{index+1}/{filename}", "w") as dest:
                    import shutil
                    shutil.copyfileobj(handle, dest)
                packed += artifact.size
        archive.writestr("results.json", json.dumps({"rows": [{"title": r["title"], "error": r.get("error", "")} for r in batch.rows], "separate_downloads": omitted}, ensure_ascii=False))
        publish()
    finally:
        archive.close()
    return {"batch_id": str(batch.id), "completed": sum(not r.get("error") for r in batch.rows), "failed": sum(bool(r.get("error")) for r in batch.rows)}
