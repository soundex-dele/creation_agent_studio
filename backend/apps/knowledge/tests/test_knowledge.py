import io

import pytest
from django.contrib.auth import get_user_model
from docx import Document
from pypdf import PdfWriter
from rest_framework.test import APIClient

from apps.enterprise.models import Membership, ProviderConfig
from apps.knowledge.extractors import DocumentExtractionError, extract_document
from apps.knowledge.execution import execute_knowledge_answer, execute_knowledge_index
from apps.knowledge.index_backend import replace_document_index
from apps.knowledge.models import KnowledgeBase, KnowledgeChunk, KnowledgeDocument
from apps.knowledge.retrieval import chunk_sections, search
from modules.execution.infrastructure.coordinator import ADAPTER_EVENT_TYPES


pytestmark = pytest.mark.django_db


def client_for(user, organization):
    client = APIClient()
    client.force_authenticate(user)
    client.credentials(HTTP_X_ORGANIZATION_ID=str(organization.id))
    return client


class Sink:
    cancelled = False

    def __init__(self):
        self.events = []

    def emit(self, event_type, payload):
        self.events.append((event_type, payload))


def test_text_extraction_and_chunking_rejects_binary():
    sections = extract_document("标题\n\n这是一段知识内容。".encode(), "source.md")
    chunks = chunk_sections(sections, 100, 10)
    assert chunks[0]["content"].startswith("标题")
    assert any("知 识 内 容" in chunk["lexical_text"] for chunk in chunks)
    with pytest.raises(DocumentExtractionError, match="NUL"):
        extract_document(b"bad\x00text", "source.txt")
    long_sections = extract_document(("知识库分块测试。" * 200).encode(), "long.txt")
    bounded = chunk_sections(long_sections, 100, 10)
    assert len(bounded) > 1
    assert all(item["token_count"] <= 100 for item in bounded)


def test_docx_extraction_and_pdf_safety_rejections():
    document = Document()
    document.add_heading("安装说明", level=1)
    document.add_paragraph("先下载客户端，再完成登录。")
    docx_bytes = io.BytesIO()
    document.save(docx_bytes)
    sections = extract_document(docx_bytes.getvalue(), "guide.docx")
    assert sections[-1].section_path == ("安装说明",)
    assert "完成登录" in sections[-1].text

    blank = PdfWriter()
    blank.add_blank_page(width=72, height=72)
    blank_bytes = io.BytesIO()
    blank.write(blank_bytes)
    with pytest.raises(DocumentExtractionError) as scanned:
        extract_document(blank_bytes.getvalue(), "scan.pdf")
    assert scanned.value.code == "no_extractable_text"

    encrypted = PdfWriter()
    encrypted.add_blank_page(width=72, height=72)
    encrypted.encrypt("secret")
    encrypted_bytes = io.BytesIO()
    encrypted.write(encrypted_bytes)
    with pytest.raises(DocumentExtractionError) as protected:
        extract_document(encrypted_bytes.getvalue(), "protected.pdf")
    assert protected.value.code == "encrypted_pdf"


def test_viewer_can_search_but_cannot_upload(tmp_path, settings):
    settings.ARTIFACT_ROOT = tmp_path
    owner = get_user_model().objects.create_user(username="knowledge-owner")
    viewer = get_user_model().objects.create_user(username="knowledge-viewer")
    organization = owner.organization_memberships.get().organization
    Membership.objects.create(
        organization=organization, user=viewer, role=Membership.Role.VIEWER)
    knowledge_base = KnowledgeBase.objects.create(
        organization=organization, name="产品资料")
    document = KnowledgeDocument.objects.create(
        organization=organization,
        knowledge_base=knowledge_base,
        created_by=owner,
        title="价格说明",
        status=KnowledgeDocument.Status.READY,
        active_revision=1,
    )
    chunk = KnowledgeChunk.objects.create(
        organization=organization,
        document=document,
        revision=1,
        position=0,
        content="高级版价格为每月一百元",
        lexical_text="高 级 版 价 格 为 每 月 一 百 元",
    )
    replace_document_index(document, [chunk])
    client = client_for(viewer, organization)
    response = client.post(
        f"/api/v1/organizations/{organization.id}/knowledge-search/",
        {"query": "价格", "knowledge_base_ids": [knowledge_base.id]},
        format="json",
    )
    assert response.status_code == 200
    assert response.data["retrieval_mode"] == "lexical"
    assert response.data["results"][0]["document_id"] == document.id
    upload = client.post(
        f"/api/v1/organizations/{organization.id}/knowledge-bases/{knowledge_base.id}/documents/",
        {"title": "blocked", "text": "content"},
        format="json",
    )
    assert upload.status_code == 403


def test_developer_upload_creates_durable_index_run(tmp_path, settings):
    settings.ARTIFACT_ROOT = tmp_path
    owner = get_user_model().objects.create_user(username="knowledge-developer")
    organization = owner.organization_memberships.get().organization
    knowledge_base = KnowledgeBase.objects.create(
        organization=organization, name="帮助中心")
    client = client_for(owner, organization)
    response = client.post(
        f"/api/v1/organizations/{organization.id}/knowledge-bases/{knowledge_base.id}/documents/",
        {"title": "使用指南", "text": "这是可以被索引的帮助文档。"},
        format="json",
    )
    assert response.status_code == 202
    document = KnowledgeDocument.objects.get(pk=response.data["document"]["id"])
    assert document.source_object_key
    assert response.data["run_url"].endswith(str(document.indexing_run_id))
    assert response.data["stream_url"].endswith(f"{document.indexing_run_id}/stream")
    assert document.indexing_run.executor_kind == "knowledge"
    assert document.indexing_run.executor_key == "knowledge-index"
    run = document.indexing_run
    sink = Sink()
    output = execute_knowledge_index({
        "run_id": str(run.id),
        "organization_id": str(organization.id),
        "owner_id": owner.id,
        "definition_snapshot": run.definition_snapshot,
        "input": run.input,
    }, sink)
    document.refresh_from_db()
    assert document.status == KnowledgeDocument.Status.READY
    assert output["chunk_count"] == 1
    assert {event_type for event_type, _payload in sink.events} <= ADAPTER_EVENT_TYPES
    mode, results = search(
        organization_id=organization.id,
        query="帮助文档",
        knowledge_base_ids=[knowledge_base.id],
        limit=10,
    )
    assert mode == "lexical"
    assert results[0]["document_id"] == document.id


def test_search_is_tenant_isolated():
    first = get_user_model().objects.create_user(username="knowledge-tenant-a")
    second = get_user_model().objects.create_user(username="knowledge-tenant-b")
    first_org = first.organization_memberships.get().organization
    second_org = second.organization_memberships.get().organization
    base = KnowledgeBase.objects.create(organization=second_org, name="秘密资料")
    document = KnowledgeDocument.objects.create(
        organization=second_org, knowledge_base=base, title="秘密",
        status=KnowledgeDocument.Status.READY, active_revision=1,
    )
    KnowledgeChunk.objects.create(
        organization=second_org, document=document, revision=1, position=0,
        content="不可见内容", lexical_text="不 可 见 内 容",
    )
    mode, results = search(
        organization_id=first_org.id, query="内容", knowledge_base_ids=[], limit=10)
    assert mode == "lexical"
    assert results == []


def test_answer_run_requires_provider_and_replays_idempotently():
    owner = get_user_model().objects.create_user(username="knowledge-answer-owner")
    organization = owner.organization_memberships.get().organization
    base = KnowledgeBase.objects.create(
        organization=organization, name="问答资料", answer_model="chat-model")
    client = client_for(owner, organization)
    url = f"/api/v1/organizations/{organization.id}/knowledge-answer-runs/"
    missing = client.post(
        url, {"query": "答案", "knowledge_base_ids": [base.id]}, format="json",
        HTTP_IDEMPOTENCY_KEY="answer-1",
    )
    assert missing.status_code == 409
    provider = ProviderConfig.objects.create(
        organization=organization,
        name="default",
        base_url="https://example.invalid/v1",
        available_models=["chat-model"],
    )
    first = client.post(
        url, {"query": "答案", "knowledge_base_ids": [base.id]}, format="json",
        HTTP_IDEMPOTENCY_KEY="answer-1",
    )
    provider.delete()
    second = client.post(
        url, {"query": "答案", "knowledge_base_ids": [base.id]}, format="json",
        HTTP_IDEMPOTENCY_KEY="answer-1",
    )
    assert first.status_code == second.status_code == 202
    assert first.data["run"]["id"] == second.data["run"]["id"]
    assert first.data["output"] is None
    assert first.data["stream_url"].endswith("/stream")
    assert second["Idempotent-Replay"] == "true"


def test_embedding_failure_degrades_to_lexical(tmp_path, settings, monkeypatch):
    settings.ARTIFACT_ROOT = tmp_path
    owner = get_user_model().objects.create_user(username="knowledge-lexical-fallback")
    organization = owner.organization_memberships.get().organization
    base = KnowledgeBase.objects.create(
        organization=organization, name="降级资料", embedding_model="embed-model")
    client = client_for(owner, organization)
    response = client.post(
        f"/api/v1/organizations/{organization.id}/knowledge-bases/{base.id}/documents/",
        {"title": "降级文档", "text": "嵌入服务异常时仍应支持关键词搜索。"},
        format="json",
    )
    document = KnowledgeDocument.objects.get(pk=response.data["document"]["id"])
    monkeypatch.setattr(
        "apps.knowledge.execution.embed_texts",
        lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("provider offline")),
    )
    run = document.indexing_run
    output = execute_knowledge_index({
        "run_id": str(run.id),
        "organization_id": str(organization.id),
        "owner_id": owner.id,
        "definition_snapshot": run.definition_snapshot,
        "input": run.input,
    }, Sink())
    document.refresh_from_db()
    assert document.status == KnowledgeDocument.Status.READY
    assert output["retrieval_mode"] == "lexical"
    assert document.metadata["embedding_warning"] == "provider offline"


def test_embedding_dimension_error_fails_without_switching_revision(
        tmp_path, settings, monkeypatch):
    settings.ARTIFACT_ROOT = tmp_path
    owner = get_user_model().objects.create_user(username="knowledge-dimension-error")
    organization = owner.organization_memberships.get().organization
    base = KnowledgeBase.objects.create(
        organization=organization, name="维度校验", embedding_model="embed-model")
    client = client_for(owner, organization)
    response = client.post(
        f"/api/v1/organizations/{organization.id}/knowledge-bases/{base.id}/documents/",
        {"title": "维度错误", "text": "这段内容不应切换为活动索引。"},
        format="json",
    )
    document = KnowledgeDocument.objects.get(pk=response.data["document"]["id"])
    monkeypatch.setattr(
        "apps.knowledge.execution.embed_texts",
        lambda *args, **kwargs: (_ for _ in ()).throw(
            RuntimeError("Embedding dimension mismatch: expected 1024, got 3.")),
    )
    run = document.indexing_run
    with pytest.raises(RuntimeError, match="dimension mismatch"):
        execute_knowledge_index({
            "run_id": str(run.id),
            "organization_id": str(organization.id),
            "owner_id": owner.id,
            "definition_snapshot": run.definition_snapshot,
            "input": run.input,
        }, Sink())
    document.refresh_from_db()
    assert document.status == KnowledgeDocument.Status.FAILED
    assert document.active_revision == 0
    assert document.error_code == "indexing_failed"


def test_answer_filters_invalid_citations(monkeypatch):
    owner = get_user_model().objects.create_user(username="knowledge-citations")
    organization = owner.organization_memberships.get().organization
    monkeypatch.setattr(
        "apps.knowledge.execution.answer_question",
        lambda *args, **kwargs: ("有效结论 [1]，无效引用 [9]。", {"total_tokens": 12}, "chat-model"),
    )
    results = [
        {
            "chunk_id": 1, "document_id": 2, "knowledge_base_id": 3,
            "knowledge_base_name": "产品资料", "title": "说明书",
            "snippet": "有效原文", "score": 0.1, "page_number": 1,
            "section_path": ["安装"], "citation": "说明书#chunk-0",
        },
        {
            "chunk_id": 4, "document_id": 5, "knowledge_base_id": 3,
            "knowledge_base_name": "产品资料", "title": "其他资料",
            "snippet": "未引用原文", "score": 0.05, "page_number": None,
            "section_path": [], "citation": "其他资料#chunk-0",
        },
    ]
    output = execute_knowledge_answer({
        "run_id": "00000000-0000-0000-0000-000000000001",
        "organization_id": str(organization.id),
        "owner_id": owner.id,
        "definition_snapshot": {
            "results": results, "retrieval_mode": "hybrid",
            "answer_provider": "default", "answer_model": "chat-model",
        },
        "input": {"query": "如何安装？"},
    }, Sink())
    assert "[9]" not in output["answer"]
    assert [item["label"] for item in output["citations"]] == ["[1]"]


def test_document_delete_is_retryable_after_storage_failure(monkeypatch):
    owner = get_user_model().objects.create_user(username="knowledge-delete-retry")
    organization = owner.organization_memberships.get().organization
    base = KnowledgeBase.objects.create(organization=organization, name="待删除资料")
    document = KnowledgeDocument.objects.create(
        organization=organization,
        knowledge_base=base,
        created_by=owner,
        title="待删除文档",
        source_object_key="knowledge/source.txt",
        status=KnowledgeDocument.Status.READY,
        active_revision=1,
    )
    client = client_for(owner, organization)
    url = (
        f"/api/v1/organizations/{organization.id}/knowledge-bases/{base.id}/"
        f"documents/{document.id}/"
    )
    monkeypatch.setattr(
        "apps.knowledge.views.delete_artifact_object",
        lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("storage offline")),
    )
    failed = client.delete(url)
    assert failed.status_code == 503
    document.refresh_from_db()
    assert document.is_deleted is True
    assert document.source_object_key

    monkeypatch.setattr("apps.knowledge.views.delete_artifact_object", lambda *args: None)
    retried = client.delete(url)
    assert retried.status_code == 204
    document.refresh_from_db()
    assert document.source_object_key == ""
